"""Build the four-slide defense deck from the UPM template (presentation.pptx).

Keeps template slides 7, 3, 4 and 5 (in that order), fills their placeholders keeping the template's formatting
and puts the figures from make_slide_figures.py and the report into the picture slots.
Run from this folder: python build_deck.py   ->  defense_slides.pptx
"""
import copy

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
TEMPLATE, OUTPUT = "presentation.pptx", "defense_slides.pptx"
FOOTER = "T. S. A. Camargo · J. Corrales Almendariz · J. Montes de Oca Martínez"

prs = Presentation(TEMPLATE)

# Keep template slides 7 (summary), 3 (tree), 4 (logistic regression) and 5 (EBM), in this order.
slide_ids = prs.slides._sldIdLst
original = list(slide_ids)
keep = [original[i - 1] for i in (7, 3, 4, 5)]
for element in original:
    slide_ids.remove(element)
    if element not in keep:
        prs.part.drop_rel(element.rId)
for element in keep:
    slide_ids.append(element)
summary, tree, logit, ebm = prs.slides


def shape(slide, shape_id):
    return next(s for s in slide.shapes if s.shape_id == shape_id)


def place(target, left, top, width, height):
    """Give a placeholder its own geometry (all four values, or it collapses to zero width)."""
    target.left, target.top, target.width, target.height = (Inches(v) for v in (left, top, width, height))


def set_text(target, paragraphs, size=None, spacing=None):
    """Replace the text, reusing the first paragraph and run as formatting templates.

    `paragraphs` is a list of paragraphs, each a string or a list of (text, bold) runs.
    """
    body = target.text_frame._txBody
    old = body.findall(f"{A}p")
    template_p = old[0]
    template_r = next(iter(template_p.iter(f"{A}r")))
    for p in old:
        body.remove(p)
    for runs in paragraphs:
        if isinstance(runs, str):
            runs = [(runs, None)]
        p = copy.deepcopy(template_p)
        for child in list(p):
            if child.tag in (f"{A}r", f"{A}br", f"{A}fld"):
                p.remove(child)
        end = p.find(f"{A}endParaRPr")
        for text, bold in runs:
            r = copy.deepcopy(template_r)
            r.find(f"{A}t").text = text
            properties = r.find(f"{A}rPr")
            if bold is not None:
                properties.set("b", "1" if bold else "0")
            if size is not None:
                properties.set("sz", str(int(size * 100)))
            if spacing is not None:
                properties.set("spc", str(spacing))
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        body.append(p)


def remove(slide, shape_id):
    element = shape(slide, shape_id)._element
    element.getparent().remove(element)


def picture(slide, path, left, top, width=None, height=None):
    kwargs = {k: Inches(v) for k, v in (("width", width), ("height", height)) if v is not None}
    return slide.shapes.add_picture(path, Inches(left), Inches(top), **kwargs)


def note(slide, text, left, top, width, height, size=9):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = text
    for p in frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.name = "Montserrat"
            r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    return box


def set_dial(group, fraction):
    """The dial is a coloured ring with a grey arc covering the unfilled part (from adj1 clockwise to 12 o'clock)."""
    arc = next(c for c in group.shapes if c._element.find(f".//{A}prstGeom").get("prst") == "blockArc")
    guides = {g.get("name"): g for g in arc._element.iter(f"{A}gd")}
    end = int(guides["adj2"].get("fmla").split()[1])
    start = (end - round((1 - fraction) * 360 * 60000)) % (360 * 60000)
    guides["adj1"].set("fmla", f"val {start}")


for slide in prs.slides:
    footer = next(s for s in slide.placeholders if s.has_text_frame and s.text_frame.text == "Título del proyecto")
    set_text(footer, [FOOTER])

# --- Slide 1: summary ------------------------------------------------------------------------------------------
set_text(shape(summary, 35), ["Summary of results"])
set_text(shape(summary, 10), ["Five-fold CV AUC of each model"])
dials = [(21, 29, 42, 43, "Decision tree", 0.752, ["17 leaves, 6 rules", "Accuracy 0.812", "F1 0.49 at t_bal"]),
         (15, 31, 44, 45, "Logistic reg.", 0.761, ["Log of money vars.", "Accuracy 0.814", "F1 0.50 at t_bal"]),
         (12, 33, 46, 47, "EBM", 0.776, ["Best in 5/5 folds", "Accuracy 0.817", "F1 0.52 at t_bal"])]
for group_id, value_id, title_id, body_id, name, auc, lines in dials:
    set_dial(shape(summary, group_id), auc)
    set_text(shape(summary, value_id), [f"{auc:.3f}"])
    set_text(shape(summary, title_id), [name], size=14)
    set_text(shape(summary, body_id), lines, size=11)
note(summary,
     "Same 80/20 train/validation split and the same five CV folds for every choice: hyperparameters, a log "
     "transform and the thresholds (t_acc for accuracy, t_bal for balanced accuracy). All models beat the majority "
     "baseline (accuracy 0.788); a black-box gradient-boosting model gets AUC 0.768. Submitted: the EBM.",
     0.75, 6.5, 8.5, 0.6)

# --- Slide 2: decision tree ------------------------------------------------------------------------------------
set_text(shape(tree, 3), ["Decision tree"])
set_text(shape(tree, 11), ["17 leaves that collapse into six rules"])
set_text(shape(tree, 7), [
    [("Flag difficulty if any rule fires ", True), ("(status = latest repayment status, ratio = payment-to-bill "
                                                    "ratio):", False)],
    [("R1 ", True), ("status ≥ 1 and ratio ≤ 0.095: risk 0.54; 15 % of clients, 39 % of all difficulties", False)],
    [("R2 ", True), ("status ≤ 0 and ratio ≤ 0.025: 0.43", False)],
    [("R3 ", True), ("status ≥ 1, ratio > 0.095, zero-payment months ≥ 4: 0.33", False)],
    [("R4–R6 ", True), ("low ratio with age ≤ 24 or with an assistance plan: 0.26–0.32", False)],
    [("Global: ", True), ("uses 7 of 15 variables; status and ratio give 85 % of its importance.", False)],
    [("Trade-off: ", True), ("easiest to apply by hand, but coarse (17 probability levels): CV AUC 0.752, the "
                             "lowest.", False)],
], size=10.5)
remove(tree, 9)
picture(tree, "../report/figures/tree_top_levels.png", 5.62, 1.3, width=3.95)
caption = shape(tree, 2)
place(caption, 5.62, 3.55, 3.95, 3.15)
set_text(caption, [
    [("Local case C (actual: difficulty)", True)],
    "Path: status ≥ 1 → ratio > 0.095 → zero-payment months ≤ 3 → ratio > 0.315: leaf risk 0.15, missed.",
    "A ratio of 4.39 on a negative bill looks safe to the tree; LR (0.50) and the EBM (0.55) flag the client.",
], size=11)

# --- Slide 3: logistic regression ------------------------------------------------------------------------------
set_text(shape(logit, 8), ["Logistic regression"])
set_text(shape(logit, 2), ["Compact and stable, but conditional and linear"])
body = shape(logit, 6)
place(body, 0.72, 3.3, 3.85, 3.75)
set_text(body, [
    [("Setup: ", True), ("standardized inputs; a CV-chosen log of the money variables adds 0.012 AUC.", False)],
    [("Global: ", True), ("risk rises with utilization, the assistance plan, delayed months and repayment "
                          "status; it falls with total payment.", False)],
    [("Caution: ", True), ("coefficients are conditional: maximum delay, credit limit and payment ratio flip "
                           "their sign (collinearity).", False)],
    [("Limit: ", True), ("straight lines only; status codes −2, −1 and 0 are forced onto a line.", False)],
    [("Local, case D: ", True), ("utilization 2.0 adds +1.80: p = 0.39, a false alarm. The tree ignores "
                                 "utilization and is right.", False)],
    [("Performance: ", True), ("CV AUC 0.761; F1 0.50 at t_bal.", False)],
], size=10.5)
remove(logit, 13)
picture(logit, "figures/lr_coefficients.png", 5.47, 1.33, width=4.25, height=2.42)
picture(logit, "figures/lr_case_d.png", 5.47, 3.92, width=4.25, height=2.42)

# --- Slide 4: EBM ----------------------------------------------------------------------------------------------
set_text(shape(ebm, 5), ["EBM"])
set_text(shape(ebm, 3), ["The best model: CV AUC 0.776"])
body = shape(ebm, 4)
place(body, 0.72, 3.0, 3.85, 4.05)
set_text(body, [
    [("Ahead of LR and the tree ", True), ("in all five CV folds; submitted for test.csv (t_acc = 0.49).", False)],
    [("Additive: ", True), ("one shape function per variable (interactions = 0), read directly from the "
                            "plots.", False)],
    [("Payment ratio is U-shaped: ", True), ("risk is lowest at 0.5–1.2 and rises again above 1.7 (catch-up "
                                             "payments, tiny bills).", False)],
    [("Zero-payment months: ", True), ("1–2 lower the risk, 3+ raise it. Forcing a monotone shape lowers the CV "
                                       "AUC (0.7756 → 0.7728), so the dip is real.", False)],
    [("Age is U-shaped: ", True), ("clients under 25 and over 57 are riskier.", False)],
    [("Assistance plan +0.41: ", True), ("a marker of earlier stress, not a causal harm.", False)],
], size=10.5)
for picture_id in (17, 18, 19, 22):
    remove(ebm, picture_id)
slot_offset = (2.81 - 2.05) / 2
picture(ebm, "figures/ebm_ratio.png", 5.11, 1.10 + slot_offset, width=2.15)
picture(ebm, "figures/ebm_zero_payment.png", 7.48, 1.10 + slot_offset, width=2.15)
picture(ebm, "figures/ebm_age.png", 7.48, 4.21 + slot_offset, width=2.15)
caption = shape(ebm, 8)
place(caption, 5.10, 4.21, 2.15, 2.81)
set_text(caption, [
    [("Local case C", True)],
    "Ratio 4.39: +0.48",
    "Delays: +0.49",
    "Plan: +0.41",
    "p = 0.55: flagged",
    "(tree: 0.15, missed)",
], size=10, spacing=0)

prs.save(OUTPUT)
print("saved", OUTPUT)
